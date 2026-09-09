# -*- coding: utf-8 -*-
"""Detect duplicate/overflow issues in output decks (QA gate)."""
from pptx import Presentation
import glob
from collections import Counter

fails = 0
for path in sorted(glob.glob("output/*.pptx")):
    prs = Presentation(path)
    titles = []
    for s in prs.slides:
        txts = [sh.text_frame.paragraphs[0].text.strip() for sh in s.shapes
                if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(txts[0] if txts else "(visual-only)")
    dupes = [t[:80] for t, c in Counter(titles).items() if c > 1 and t != "(visual-only)"]
    print(f"{path}: {len(prs.slides._sldIdLst)} slides" + (f"  DUPES={dupes}" if dupes else "  OK-no-dupes"))
    if dupes:
        fails += 1
    # overflow heuristic: too many paragraphs in one textbox
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and len(sh.text_frame.paragraphs) > 12:
                print(f"  WARN {path} s{i}: {len(sh.text_frame.paragraphs)} paras — may overflow")
print("FAIL" if fails else "PASS")