# -*- coding: utf-8 -*-
"""Module02 — Simple visualizations deck (native shapes, light theme, RTL)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x33, 0x33, 0x33)
TITLE_C = RGBColor(0x1F, 0x3B, 0x5C)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x66, 0x66, 0x66)
FONT_AR = "Arial"
FONT_EN = "Calibri"
FILL_A = RGBColor(0xDC, 0xE6, 0xF1)   # light blue
FILL_B = RGBColor(0xE2, 0xEF, 0xDA)   # light green
FILL_C = RGBColor(0xFD, 0xE9, 0xD9)   # light orange
FILL_D = RGBColor(0xED, 0xED, 0xED)   # light gray

SW, SH = Inches(13.333), Inches(7.5)


def set_rtl(par, rtl=True):
    pPr = par._p.get_or_add_pPr()
    pPr.set("rtl", "1" if rtl else "0")


def prs_new():
    p = Presentation()
    p.slide_width, p.slide_height = SW, SH
    return p


def slide_new(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    tb = s.shapes.add_textbox(Inches(0.7), Inches(0.15), Inches(12.0),
                              Inches(0.7))
    par = tb.text_frame.paragraphs[0]
    par.alignment = PP_ALIGN.RIGHT
    set_rtl(par)
    r = par.add_run(); r.text = title
    r.font.size, r.font.bold, r.font.name = Pt(30), True, FONT_AR
    r.font.color.rgb = TITLE_C
    return s


def card(slide, x, y, w, h, ar, en=None, fill=FILL_A, size=16, sub=12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = ACCENT; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER; set_rtl(p1)
    r1 = p1.add_run(); r1.text = ar
    r1.font.size, r1.font.bold, r1.font.color.rgb = Pt(size), True, INK
    r1.font.name = FONT_AR
    if en:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; set_rtl(p2, False)
        r2 = p2.add_run(); r2.text = en
        r2.font.size, r2.font.color.rgb = Pt(sub), MUTED
        r2.font.name = FONT_EN
    return sh


def arrow(slide, x, y, w=Inches(0.5), h=Inches(0.3), left=True):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.LEFT_ARROW if left else MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = ACCENT
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def note(slide, ar, en=None, y=Inches(6.7)):
    tb = slide.shapes.add_textbox(Inches(0.7), y, Inches(12.0), Inches(0.6))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; set_rtl(p)
    r = p.add_run(); r.text = ar
    r.font.size, r.font.color.rgb = Pt(14), MUTED; r.font.name = FONT_AR
    if en:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.RIGHT; set_rtl(p2, False)
        r2 = p2.add_run(); r2.text = en
def v01_roadmap(prs):
    s = slide_new(prs, "خريطة المسار / Track Roadmap")
    steps = [("البيانات", "Data"), ("التخزين", "Storage"),
             ("الخصائص", "Features"), ("الفهارس المتجهية", "Vector Indexes"),
             ("التصنيف", "Labeling"), ("الحوكمة", "Governance"),
             ("MLOps", "MLOps"), ("المختبر", "Lab")]
    w, gap, x = Inches(1.35), Inches(0.16), Inches(0.55)
    for i, (ar, en) in enumerate(steps):
        fill = [FILL_A, FILL_B, FILL_C, FILL_D][i % 4]
        card(s, x, Inches(2.6), w, Inches(1.5), ar, en, fill, size=14, sub=10)
        if i < len(steps) - 1:
            arrow(s, x + w + Inches(0.01), Inches(3.2), Inches(0.14), Inches(0.3))
        x += w + gap
    note(s, "رحلة الوحدة الثانية من البيانات الخام إلى نقطة نهاية RAG.",
         "Module 02 journey: from raw data to a RAG endpoint.")


def v02_medallion(prs):
    s = slide_new(prs, "المعمارية المتدرِّجة Medallion / Medallion Layers")
    rows = [("Bronze — بيانات خام غير قابلة للتغيير", "Raw immutable data", FILL_C),
            ("Silver — بيانات منقَّاة ومنظَّمة", "Cleansed & structured data", FILL_D),
            ("Gold — تجميعات جاهزة للتحليل", "Analytics-ready aggregates", FILL_B)]
    y = Inches(1.3)
    for ar, en, fill in rows:
        card(s, Inches(2.5), y, Inches(8.3), Inches(1.35), ar, en, fill, size=18, sub=13)
        y += Inches(1.65)
    arrow(s, Inches(6.45), Inches(2.68), Inches(0.4), Inches(0.3))
    arrow(s, Inches(6.45), Inches(4.33), Inches(0.4), Inches(0.3))
    note(s, "الحوكمة عبر تحسين الجودة بين الطبقات — إدارة مركزية.",
         "Governance via quality improvement between layers — central ownership.")


def v03_patterns(prs):
    s = slide_new(prs, "أنماط معالجة البيانات الثلاثة / Three Data Patterns")
    cols = [("المعمارية المتدرِّجة", "Medallion", "فريق مركزي — تحليلات وAI/ML/RAG", FILL_A),
            ("مستودع البيانات", "Data Warehouse", "فرق BI — تقارير وذكاء أعمال", FILL_B),
            ("شبكة البيانات", "Data Mesh", "فرق المجالات — مؤسسات كبيرة", FILL_C)]
    x = Inches(0.7)
    for ar, en, use, fill in cols:
        card(s, x, Inches(1.5), Inches(3.9), Inches(1.0), ar, en, fill, size=18, sub=13)
        card(s, x, Inches(2.7), Inches(3.9), Inches(2.8), use, None, FILL_D, size=15)
        x += Inches(4.1)
    note(s, "المعايير: الفكرة الجوهرية / المالك / الحوكمة / حالة الاستخدام المثلى.",
         "Criteria: core idea / ownership / governance / best use case.")


def v04_contract(prs):
    s = slide_new(prs, "منتج البيانات وعقد البيانات / Data Product & Contract")
    card(s, Inches(0.9), Inches(1.8), Inches(4.6), Inches(3.4),
         "منتج البيانات", "Data Product — موثوق • موثّق • قابل للاكتشاف", FILL_A, size=20, sub=14)
    card(s, Inches(7.9), Inches(1.8), Inches(4.6), Inches(1.0),
         "عقد البيانات", "Data Contract", FILL_B, size=20, sub=14)
    for i, (ar, en) in enumerate([("المخطط", "Schema"), ("الأمان", "Security"),
                                  ("مستوى الخدمة", "SLA")]):
        card(s, Inches(8.3), Inches(3.0 + i * Inches(0.8)), Inches(3.8),
             Inches(0.65), f"{ar} / {en}", None, FILL_D, size=14)
    arrow(s, Inches(5.75), Inches(3.3), Inches(1.9), Inches(0.4))
    note(s, "وحدة التسليم الحديثة في هندسة البيانات.",
         "The modern delivery unit in data engineering.")


def v05_rag_pipeline(prs):
    s = slide_new(prs, "خط أنابيب RAG / RAG Pipeline")
    steps = [("1) الفهرسة", "Indexing"), ("2) الاسترجاع", "Retrieval"),
             ("3) التوليد", "Generation"), ("4) التقييم", "Evaluation"),
             ("5) التحسين", "Optimization")]
    w, gap, x = Inches(2.15), Inches(0.28), Inches(0.6)
    for i, (ar, en) in enumerate(steps):
        fill = [FILL_A, FILL_B, FILL_C, FILL_D, FILL_A][i]
        card(s, x, Inches(2.2), w, Inches(1.4), ar, en, fill, size=16, sub=12)
        if i < 4:
            arrow(s, x + w + Inches(0.02), Inches(2.75), Inches(0.24), Inches(0.3))
        x += w + gap
    detail = ["تقسيم تكراري + HNSW", "بحث هجين + إعادة ترتيب",
              "موجِّه مقيَّد بالمصادر", "مثلث RAG (RAGAS)",
              "تخزين مؤقت + توجيه"]
    x = Inches(0.6)
    for d in detail:
        card(s, x, Inches(4.1), Inches(2.15), Inches(1.0), d, None, FILL_D, size=12)
        x += Inches(2.43)
    note(s, "من الفهرسة إلى التقييم — المسار التطبيقي في مختبر L03.",
         "From indexing to evaluation — applied in lab L03.")


def v06_chunking(prs):
    s = slide_new(prs, "استراتيجيات تقسيم النصوص / Chunking Strategies")
    rows = [("ثابت الحجم", "Fixed-size", "أبسطها — يقطع الجُمل", FILL_D),
            ("دلالي", "Semantic", "تماسك أفضل — حساس للعتبة", FILL_D),
            ("تكراري ★", "Recursive", "الموصى به افتراضيًا", FILL_B),
            ("بحسب البنية", "Structure-based", "عناوين وأقسام — يُدمج مع التكراري", FILL_D),
            ("بالنموذج اللغوي", "LLM-based", "الأدق — الأغلى حسابيًا", FILL_D)]
    y = Inches(1.35)
    for ar, en, d, fill in rows:
        card(s, Inches(0.9), y, Inches(2.8), Inches(0.85), ar, en, fill, size=15, sub=11)
        card(s, Inches(3.9), y, Inches(6.2), Inches(0.85), d, None, FILL_D, size=13)
        y += Inches(1.0)
    note(s, "★ التقسيم التكراري: فصل بالفقرات ثم تقسيم ما تجاوز الحد الأقصى.",
         "★ Recursive: split by paragraphs, then oversized chunks.")


def v07_triad(prs):
    s = slide_new(prs, "مثلث تقييم RAG / RAG Evaluation Triad")
    tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                             Inches(4.1), Inches(1.6), Inches(5.0), Inches(4.2))
    tri.fill.solid(); tri.fill.fore_color.rgb = FILL_A
    tri.line.color.rgb = ACCENT; tri.shadow.inherit = False
    card(s, Inches(4.9), Inches(1.3), Inches(3.4), Inches(0.8),
         "الوفاء بالمصدر", "Faithfulness", FILL_B, size=16, sub=12)
    card(s, Inches(2.4), Inches(5.3), Inches(3.4), Inches(0.8),
         "ملاءمة الإجابة", "Answer Relevancy", FILL_C, size=16, sub=12)
    card(s, Inches(7.4), Inches(5.3), Inches(3.4), Inches(0.8),
         "دقة السياق", "Context Precision", FILL_D, size=16, sub=12)
    note(s, "يُقاس بإطار RAGAS — الوفاء يمنع الهلوسة والملاءمة تجعل الإجابة نافعة.",
         "Measured with RAGAS — faithfulness prevents hallucination; relevancy keeps answers useful.")


def v13_grounded_answer(prs):
    """Show the learner why evidence and citations make a RAG answer trustworthy."""
    s = slide_new(prs, "نمط الإجابة المؤرَّضة / Grounded Answer Pattern")
    card(s, Inches(0.7), Inches(1.55), Inches(3.0), Inches(1.2),
         "سؤال المتعلِّم", "Learner question", FILL_A, size=18, sub=13)
    card(s, Inches(0.7), Inches(2.95), Inches(3.0), Inches(1.25),
         "ما الاستراتيجية الافتراضية للتقسيم؟",
         "Which chunking default should I use?", FILL_D, size=14, sub=11)
    arrow(s, Inches(3.85), Inches(3.35), Inches(0.65), Inches(0.3), left=False)

    card(s, Inches(4.65), Inches(1.4), Inches(3.35), Inches(1.1),
         "أدلة مسترجعة", "Retrieved evidence", FILL_B, size=18, sub=13)
    card(s, Inches(4.65), Inches(2.75), Inches(3.35), Inches(0.72),
         "[1] الفقرات ← الجمل ← التقسيم", "[1] paragraphs → sentences → split", FILL_D, size=12, sub=10)
    card(s, Inches(4.65), Inches(3.62), Inches(3.35), Inches(0.72),
         "[2] التكراري هو الخيار الموصى به", "[2] recursive is the recommended default", FILL_D, size=12, sub=10)
    arrow(s, Inches(8.15), Inches(3.35), Inches(0.65), Inches(0.3), left=False)

    card(s, Inches(8.95), Inches(1.55), Inches(3.65), Inches(1.2),
         "إجابة موثقة", "Grounded answer", FILL_C, size=18, sub=13)
    card(s, Inches(8.95), Inches(2.95), Inches(3.65), Inches(1.25),
         "ابدأ بالتقسيم التكراري [1][2]",
         "Start with recursive chunking [1][2]", FILL_D, size=14, sub=11)
    card(s, Inches(2.4), Inches(5.35), Inches(8.55), Inches(0.75),
         "لا دليل مسترجع؟ قل: لا أعرف / No retrieved evidence? Say: I don't know.",
         None, FILL_C, size=14)
    note(s, "القاعدة: لا تتجاوز الأدلة، واربط كل ادعاء بمصدر.",
         "Rule: do not go beyond the evidence; connect every claim to a source.")


def v08_genai_arch(prs):
    s = slide_new(prs, "المعمارية المرجعية للذكاء التوليدي / GenAI Reference Architecture")
    layers = [("أ. بوابة منصة GenAI", "A. Platform Portal — POC → MVP → PROD", FILL_A),
              ("ب. الأتمتة والامتثال", "B. Automation & Compliance — CI/CD gates", FILL_B),
              ("ج. الخدمات المشتركة", "C. Shared Services — Gateway, Prompts, Audit, Guardrails", FILL_C),
              ("د. الحوكمة والمراقبة", "D. Governance & Monitoring — Responsible AI, FinOps", FILL_D)]
    y = Inches(1.4)
    for ar, en, fill in layers:
        card(s, Inches(1.2), y, Inches(10.9), Inches(1.15), ar, en, fill, size=18, sub=13)
        y += Inches(1.32)
    note(s, "النمط المرجعي المؤسسي (McKinsey) — نافذة موحدة فوق خدمات مشتركة محكومة.",
         "Enterprise reference pattern (McKinsey) — a unified pane over governed shared services.")


def v09_cost(prs):
    s = slide_new(prs, "التكلفة عمليًا: التضمين مقابل التوليد / Cost: Embedding vs Generation")
    card(s, Inches(2.0), Inches(4.2), Inches(2.4), Inches(1.6),
         "التضمين 1x", "Embedding — cheap", FILL_B, size=18, sub=12)
    card(s, Inches(5.2), Inches(1.7), Inches(2.4), Inches(4.1),
         "التوليد ~10x", "Generation — costly", FILL_C, size=18, sub=12)
    card(s, Inches(8.6), Inches(3.4), Inches(2.9), Inches(2.4),
         "الحل: تخزين مؤقت + توجيه", "Fix: caching + routing", FILL_A, size=15, sub=12)
    note(s, "قاعدة FinOps: قِس التكلفة لكل طلب قبل التوسع.",
         "FinOps rule: measure per-request cost before scaling.")


def v10_concepts(prs):
    s = slide_new(prs, "من LLM إلى الأنظمة الوكيلة / From LLM to Agentic AI")
    steps = [("LLM", "تنبؤ بالرموز التالية", "Predicts next tokens"),
             ("الذكاء التوليدي", "يولّد محتوى جديدًا", "Generates new content"),
             ("الوكلاء Agents", "نية + استدعاء أدوات", "Intent + tool calls"),
             ("الأنظمة الوكيلة", "تخطيط وتنفيذ متعدد الخطوات", "Multi-step planning")]
    w, gap, x = Inches(2.6), Inches(0.5), Inches(0.7)
    for i, (ar, d, den) in enumerate(steps):
        fill = [FILL_D, FILL_A, FILL_C, FILL_B][i]
        card(s, x, Inches(1.8), w, Inches(1.0), ar, None, fill, size=18)
        card(s, x, Inches(3.0), w, Inches(1.6), d, den, FILL_D, size=13, sub=11)
        if i < 3:
            arrow(s, x + w + Inches(0.06), Inches(2.15), Inches(0.38), Inches(0.3))
        x += w + gap
    note(s, "Agentic RAG = بحث متجهي + جلب عبر الواجهات بذاكرة متجددة وتحقق قبل التسليم.",
         "Agentic RAG = vector search + API fetching with refreshed memory and verification.")


def v11_feature_store(prs):
    s = slide_new(prs, "معمارية مخزن الخصائص / Feature Store Architecture")
    # source
    card(s, Inches(0.7), Inches(2.8), Inches(2.2), Inches(1.3),
         "مصادر البيانات", "Data Sources", FILL_D, size=14, sub=11)
    arrow(s, Inches(2.98), Inches(3.3), Inches(0.45), Inches(0.3))
    # pipeline
    card(s, Inches(3.5), Inches(2.8), Inches(2.4), Inches(1.3),
         "خط الخصائص", "Feature Pipeline\nهندسة + تحويل", FILL_A, size=14, sub=11)
    # two stores
    arrow(s, Inches(5.98), Inches(2.55), Inches(0.45), Inches(0.3))
    arrow(s, Inches(5.98), Inches(3.95), Inches(0.45), Inches(0.3))
    card(s, Inches(6.5), Inches(1.7), Inches(2.7), Inches(1.3),
         "المخزن غير المتصل", "Offline Store\nتدريب النماذج", FILL_B, size=13, sub=11)
    card(s, Inches(6.5), Inches(3.9), Inches(2.7), Inches(1.3),
         "المخزن المتصل", "Online Store\nاستدلال منخفض زمن الاستجابة", FILL_C, size=13, sub=11)
    # consumers
    card(s, Inches(10.0), Inches(1.7), Inches(2.4), Inches(1.3),
         "التدريب", "Training", FILL_D, size=14, sub=11)
    card(s, Inches(10.0), Inches(3.9), Inches(2.4), Inches(1.3),
         "الاستدلال", "Serving", FILL_D, size=14, sub=11)
    arrow(s, Inches(9.24), Inches(2.2), Inches(0.72), Inches(0.3))
    arrow(s, Inches(9.24), Inches(4.4), Inches(0.72), Inches(0.3))
    note(s, "قيمة مخزن الخصائص: نفس تعريف الخصائص للتدريب والاستدلال — لا انحراف.",
         "Core value: one feature definition for training and serving — no skew.")


def v12_labeling_flow(prs):
    s = slide_new(prs, "سير عمل التصنيف بمساعدة LLM / LLM-assisted Labeling Workflow")
    steps = [("بيانات خام", "Raw data", FILL_D),
             ("تسمية مبدئية بالـ LLM", "LLM pre-labeling", FILL_A),
             ("مراجعة وتصحيح بشري", "Human review (HITL)", FILL_B),
             ("تسميات معتمدة", "Approved labels", FILL_C),
             ("تدريب النموذج", "Model training", FILL_D)]
    w, gap, x = Inches(2.15), Inches(0.35), Inches(0.7)
    for i, (ar, en, fill) in enumerate(steps):
        card(s, x, Inches(2.3), w, Inches(1.4), ar, en, fill, size=15, sub=12)
        if i < len(steps) - 1:
            arrow(s, x + w + Inches(0.02), Inches(2.85), Inches(0.3), Inches(0.3))
        x += w + gap
    # feedback loop
    card(s, Inches(3.0), Inches(4.4), Inches(7.4), Inches(0.8),
         "حلقة تغذية راجعة: تصحيحات البشر تُحسِّن موجِّهات الـ LLM",
         "Feedback loop: human corrections improve LLM prompts",
         FILL_D, size=13, sub=11)
    note(s, "الـ LLM يسرّع التسمية والبشر يضمنون الجودة — أفضل الممارسات الحديثة.",
         "LLM accelerates labeling; humans guarantee quality — modern best practice.")


if __name__ == "__main__":
    prs = prs_new()
    v01_roadmap(prs); v02_medallion(prs); v03_patterns(prs)
    v04_contract(prs); v05_rag_pipeline(prs); v06_chunking(prs)
    v07_triad(prs); v08_genai_arch(prs); v09_cost(prs); v10_concepts(prs)
    v11_feature_store(prs); v12_labeling_flow(prs); v13_grounded_answer(prs)
    out = "output/Module02_Visualizations.pptx"
    prs.save(out)
    print(out, len(prs.slides._sldIdLst), "slides")
