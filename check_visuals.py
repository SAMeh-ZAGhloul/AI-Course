from pptx import Presentation

visuals = {
    'v01_roadmap': 'Track Roadmap',
    'v02_medallion': 'Medallion Layers',
    'v03_patterns': 'Three Data Patterns',
    'v04_contract': 'Data Product & Contract',
    'v05_rag_pipeline': 'RAG Pipeline',
    'v06_chunking': 'Chunking Strategies',
    'v07_triad': 'RAG Evaluation Triad',
    'v08_genai_arch': 'GenAI Reference Architecture',
    'v09_cost': 'Embedding vs Generation',
    'v10_concepts': 'From LLM to Agentic AI',
    'v11_feature_store': 'Feature Store Architecture',
    'v12_labeling_flow': 'LLM-assisted Labeling Workflow',
    'v13_grounded_answer': 'Grounded Answer Pattern',
}
found = {k: [] for k in visuals}
for v in ['V06', 'V07', 'V08', 'V09', 'V10', 'V11', 'V12']:
    p = Presentation(f'output/Module02_{v}_Updated.pptx')
    for i, s in enumerate(p.slides, 1):
        txt = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        for k, m in visuals.items():
            if m in txt:
                found[k].append(f'{v}:{i}')
for k, locs in found.items():
    print('OK  ' if locs else 'MISS', k, '->', ', '.join(locs))
p = Presentation('output/Module02_Visualizations.pptx')
print('standalone Module02_Visualizations.pptx:', len(p.slides._sldIdLst), 'slides')
