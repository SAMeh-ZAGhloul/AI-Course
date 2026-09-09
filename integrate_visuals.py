# -*- coding: utf-8 -*-
"""Insert visualization diagram slides into the updated module decks."""
from pptx import Presentation
import generate_visuals as gv
import generate_decks as gd
from deck_data import DECKS


def slide_has(prs, marker):
    for s in prs.slides:
        txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if marker in txt:
            return True
    return False


VISUAL_TITLES = {
    "v01_roadmap": "خريطة المسار / Track Roadmap",
    "v02_medallion": "المعمارية المتدرِّجة Medallion / Medallion Layers",
    "v03_patterns": "أنماط معالجة البيانات الثلاثة / Three Data Patterns",
    "v04_contract": "منتج البيانات وعقد البيانات / Data Product & Contract",
    "v05_rag_pipeline": "خط أنابيب RAG / RAG Pipeline",
    "v06_chunking": "استراتيجيات تقسيم النصوص / Chunking Strategies",
    "v07_triad": "مثلث تقييم RAG / RAG Evaluation Triad",
    "v08_genai_arch": "المعمارية المرجعية للذكاء التوليدي / GenAI Reference Architecture",
    "v09_cost": "التكلفة عمليًا: التضمين مقابل التوليد / Cost: Embedding vs Generation",
    "v10_concepts": "من LLM إلى الأنظمة الوكيلة / From LLM to Agentic AI",
    "v11_feature_store": "معمارية مخزن الخصائص / Feature Store Architecture",
    "v12_labeling_flow": "سير عمل التصنيف بمساعدة LLM / LLM-assisted Labeling Workflow",
    "modules_overview": "وحدات المقرر / Course Modules",
}


def insert(prs, draw_fn, after_title):
    """Insert visual after the slide whose title contains after_title (idempotent).

    Pre-checks the exact visual title — if present, skips drawing entirely
    (no orphan slide parts, no zip duplicate warnings on re-run).
    """
    want = VISUAL_TITLES.get(draw_fn.__name__)
    if want:
        for s in prs.slides:
            txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            if want in txt:
                return False
    idx = None
    for i, s in enumerate(prs.slides):
        txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if after_title in txt:
            idx = i
    n_before = len(prs.slides._sldIdLst)
    draw_fn(prs)
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert((idx + 1) if idx is not None else n_before, el)
    return True


def modules_overview(prs):
    """Course-modules description slide (V06–V12), one line per module (no overflow)."""
    pairs = [(f"{code} — {DECKS[code]['title_ar']} / {DECKS[code]['title_en']}", None)
             for code in sorted(DECKS)]
    gd.content_slide(prs, "وحدات المقرر / Course Modules", pairs, size=16)


# one visual per concept — deduped (P0 fix); anchor = title substring
jobs = {
    "output/Module02_V06_Updated.pptx": [
        (gv.v01_roadmap, "أجندة العرض"),
        (modules_overview, "Track Roadmap"),
    ],
    "output/Module02_V07_Updated.pptx": [
        (gv.v02_medallion, "معمارية Lakehouse"),
        (gv.v04_contract, "منتجات البيانات وعقودها"),
        (gv.v03_patterns, "متى تختار كل نموذج"),
    ],
    "output/Module02_V08_Updated.pptx": [
        (gv.v11_feature_store, "اتساق Offline/Online"),
    ],
    "output/Module02_V09_Updated.pptx": [
        (gv.v05_rag_pipeline, "قواعد البيانات المتجهية"),
        (gv.v06_chunking, "استراتيجيات التقسيم"),
        (gv.v07_triad, "التقييم + 5) التحسين"),
    ],
    "output/Module02_V10_Updated.pptx": [
        (gv.v12_labeling_flow, "التصنيف بمساعدة LLM"),
    ],
    "output/Module02_V11_Updated.pptx": [
        (gv.v08_genai_arch, "المعمارية المرجعية للذكاء التوليدي (2/2)"),
    ],
    "output/Module02_V12_Updated.pptx": [
        (gv.v10_concepts, "تمييز المفاهيم"),
        (gv.v09_cost, "التكلفة والأداء عمليًا"),
    ],
}

for path, inserts in jobs.items():
    prs = Presentation(path)
    for fn, anchor in inserts:
        insert(prs, fn, anchor)
    # renumber footers sequentially after inserts
    code = path.split("_V")[1][:3]
    total = len(prs.slides._sldIdLst)
    for i, s in enumerate(prs.slides, 1):
        gd.footer(s, code, i, total)
    tmp = path + ".tmp"
    prs.save(tmp)
    import os
    os.replace(tmp, path)
    print(path, "->", len(prs.slides._sldIdLst), "slides")
