"""Cross-deck regression checks for presentable Module 02 decks."""
from pathlib import Path
import re
import unittest

from pptx import Presentation


OUTPUT = Path("output")
FOOTER_TEXT = "Module 02 — Data & AI Engineering Track"
PAGE = re.compile(r"^V\d{2}\s+•\s+\d+/\d+$")


class OutputDeckTests(unittest.TestCase):
    def test_updated_module_decks_have_one_valid_footer_per_slide(self):
        for path in sorted(OUTPUT.glob("Module02_V??_Updated.pptx")):
            presentation = Presentation(path)
            for number, slide in enumerate(presentation.slides, 1):
                footer_count = sum(
                    1
                    for shape in slide.shapes
                    if shape.has_text_frame and FOOTER_TEXT in shape.text_frame.text
                )
                pages = [
                    shape.text_frame.text.strip()
                    for shape in slide.shapes
                    if shape.has_text_frame and "•" in shape.text_frame.text
                ]
                self.assertEqual(footer_count, 1, f"{path.name} slide {number}")
                self.assertEqual(sum(bool(PAGE.fullmatch(page)) for page in pages), 1,
                                 f"{path.name} slide {number}: {pages}")


if __name__ == "__main__":
    unittest.main()
