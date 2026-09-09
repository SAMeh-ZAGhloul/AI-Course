"""Regression checks for the generated V09 bilingual teaching deck."""
from collections import Counter
from pathlib import Path
import unittest

from pptx import Presentation


DECK = Path("output/Module02_V09_Updated.pptx")


def slide_text(slide):
    return " ".join(
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    )


class V09DeckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.presentation = Presentation(DECK)
        cls.texts = [slide_text(slide) for slide in cls.presentation.slides]

    def test_has_the_complete_visual_learning_flow(self):
        expected_titles = (
            "RAG Pipeline",
            "Chunking Strategies",
            "RAG Evaluation Triad",
            "Grounded Answer Pattern",
        )
        deck_text = "\n".join(self.texts)
        for title in expected_titles:
            self.assertIn(title, deck_text)

    def test_keeps_every_slide_bilingual(self):
        for number, text in enumerate(self.texts, 1):
            self.assertRegex(text, r"[\u0600-\u06ff]", f"slide {number} lacks Arabic")
            self.assertRegex(text, r"[A-Za-z]", f"slide {number} lacks English")

    def test_has_no_duplicate_titles(self):
        titles = []
        for slide in self.presentation.slides:
            text = slide_text(slide)
            titles.append(text.split("\n", 1)[0].strip())
        duplicates = [title for title, count in Counter(titles).items() if count > 1]
        self.assertEqual(duplicates, [])


if __name__ == "__main__":
    unittest.main()
